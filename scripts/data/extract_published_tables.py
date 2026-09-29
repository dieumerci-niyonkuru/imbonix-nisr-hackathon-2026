"""Turn downloaded NISR Excel tables into tidy district and sector datasets.

Run after scripts/data/download_nisr_public.py:

    pip install -r scripts/data/requirements.txt
    python scripts/data/extract_published_tables.py

Outputs (small published aggregates, safe to commit) in data/extracts/:
  district_indicators_long.csv   one row per district x indicator, with SE / CI where NISR publishes them
  district_indicators_wide.csv   one row per district, one column per indicator value
  sector_census2022.csv          416 sectors: census non-monetary poverty, MPI and population
  lfs_district_2017_2025.csv     LFS district labour indicators by year
  timeline_national.csv          Rwanda over the years: census totals, EICV, DHS and LFS rounds, province poverty
  district_population_2023_2032.csv  each district's projected population, year by year
"""

import csv
import re
from collections import defaultdict
from pathlib import Path

import openpyxl
import xlrd

REPO = Path(__file__).resolve().parents[2]
RAW = REPO / "data" / "raw" / "nisr" / "published"
OUT = REPO / "data" / "extracts"

PROVINCE = {
    "Nyarugenge": "Kigali City", "Gasabo": "Kigali City", "Kicukiro": "Kigali City",
    "Nyanza": "South", "Gisagara": "South", "Nyaruguru": "South", "Huye": "South",
    "Nyamagabe": "South", "Ruhango": "South", "Muhanga": "South", "Kamonyi": "South",
    "Karongi": "West", "Rutsiro": "West", "Rubavu": "West", "Nyabihu": "West",
    "Ngororero": "West", "Rusizi": "West", "Nyamasheke": "West",
    "Rulindo": "North", "Gakenke": "North", "Musanze": "North", "Burera": "North", "Gicumbi": "North",
    "Rwamagana": "East", "Nyagatare": "East", "Gatsibo": "East", "Kayonza": "East",
    "Kirehe": "East", "Ngoma": "East", "Bugesera": "East",
}
DISTRICTS = list(PROVINCE)
# Spellings found in the published workbooks: "goma" (RPHC5 Table 59), "Nyahibu" (LFS 2025 Table 23).
ALIASES = {"goma": "Ngoma", "nyahibu": "Nyabihu"}


def district_name(value):
    """Return the canonical district name for a cell, or None."""
    if not isinstance(value, str):
        return None
    text = re.sub(r"\s+", " ", value).strip().rstrip(":").lower()
    for name in DISTRICTS:
        if text == name.lower():
            return name
    return ALIASES.get(text)


def number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def xlsx_rows(path, sheet):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
    workbook.close()
    return rows


def xls_rows(path, sheet):
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    sheet_obj = workbook.sheet_by_name(sheet)
    return [sheet_obj.row_values(r) for r in range(sheet_obj.nrows)]


records = []


def add(district, indicator_id, indicator, value, unit, year, source, table, status="observed",
        se=None, ci_low=None, ci_high=None):
    if value is None:
        return
    records.append({
        "district": district, "province": PROVINCE[district], "indicator_id": indicator_id,
        "indicator": indicator, "value": round(value, 4), "se": se, "ci_low": ci_low, "ci_high": ci_high,
        "unit": unit, "year": year, "source": source, "table": table, "status": status,
    })


# --- EICV7 main-indicator annex tables: label | estimate | SE | CI low | CI high | CV | deff
EICV7_FILE = RAW / "eicv7" / "EICV7_Tables_MainIndicatorReport.xlsx"
# (Tables A1.3, A2.4, A3.5, A3.6 and A4.8 stop at province level, so they are not listed here.)
EICV7_TABLES = {
    "Table A9.23": ("eicv7_poverty_rate", "Poverty headcount (% of people)", "%"),
    "Table A3.7": ("eicv7_health_insurance", "Has health insurance (% of people)", "%"),
    "Table A5.12": ("eicv7_umudugudu", "Households living in umudugudu (%)", "%"),
    "Table A5.13": ("eicv7_metal_roof", "Households with metal-sheet roof (%)", "%"),
    "Table A5.14": ("eicv7_electricity_lighting", "Households lit mainly by electricity (%)", "%"),
    "Table A5.15": ("eicv7_internet_home", "Households with internet access at home (%)", "%"),
    "Table A5.16": ("eicv7_improved_water", "Households using improved drinking water (%)", "%"),
    "Table A5.17": ("eicv7_improved_sanitation", "Households using improved sanitation (%)", "%"),
    "Table A5.18": ("eicv7_hh_mobile_phone", "Households owning any mobile phone (%)", "%"),
    "Table A5.19": ("eicv7_hh_smartphone", "Households owning a smartphone (%)", "%"),
    "Table A6.20": ("eicv7_workforce_ratio", "Workforce-to-population ratio (%)", "%"),
    "Table A8.22": ("eicv7_hh_sending_transfers", "Households sending transfers to other households (%)", "%"),
}
for sheet, (indicator_id, label, unit) in EICV7_TABLES.items():
    rows = xlsx_rows(EICV7_FILE, sheet)
    in_district_block = not any(isinstance(r[0], str) and r[0].strip().lower().startswith("district") for r in rows)
    seen = set()
    for row in rows:
        if isinstance(row[0], str) and row[0].strip().lower().startswith("district"):
            in_district_block = True
            continue
        name = district_name(row[0])
        if in_district_block and name and name not in seen:
            seen.add(name)
            add(name, indicator_id, label, number(row[1]), unit, "2023/24", "NISR EICV7 Main Indicators Report (tables)",
                sheet, se=number(row[2]), ci_low=number(row[3]), ci_high=number(row[4]))
    missing = set(DISTRICTS) - seen
    if missing:
        print(f"WARNING {sheet}: missing {sorted(missing)}")

# --- FinScope 2024 district access strand (transcribed from report Figure 13; see extracts README)
for row in csv.DictReader(open(OUT / "finscope2024_district_access_strand.csv", encoding="utf-8")):
    name = district_name(row["district"])
    if not name:
        continue
    source, table = "NISR/AFR FinScope 2024 report", "Figure 13 (district access strand)"
    for column, indicator_id, label in [
        ("banked_pct", "finscope_banked", "Banked adults (%)"),
        ("other_formal_nonbank_pct", "finscope_other_formal_only", "Formal non-bank only, e.g. mobile money / SACCO (%)"),
        ("informal_only_pct", "finscope_informal_only", "Informal only (%)"),
        ("excluded_pct", "finscope_excluded", "Financially excluded adults (%)"),
    ]:
        add(name, indicator_id, label, number(row[column]), "% of adults 16+", "2024", source, table)
    banked, other = number(row["banked_pct"]), number(row["other_formal_nonbank_pct"])
    if banked is not None and other is not None:
        add(name, "finscope_formally_included", "Formally included adults (banked + other formal) (%)",
            banked + other, "% of adults 16+", "2024", source, table, status="calculated")
        add(name, "finscope_not_formally_included", "Adults not formally included: informal only or excluded (%)",
            100 - banked - other, "% of adults 16+", "2024", source, f"{table}; 100 minus banked minus other formal",
            status="calculated")

# --- Census 2022: non-monetary poverty and MPI by district (full count)
NMP = RAW / "census2022" / "Non Monetary Poverty Thematic Report.xls"
for row in xls_rows(NMP, "Table4_2"):
    name = district_name(row[0])
    if name:
        for column, indicator_id, label in [
            (1, "census_nonpoor", "Non-poor, census non-monetary measure (% of people)"),
            (2, "census_vulnerable", "Vulnerable, census non-monetary measure (% of people)"),
            (3, "census_moderately_poor", "Moderately poor, census non-monetary measure (% of people)"),
            (4, "census_severely_poor", "Severely poor, census non-monetary measure (% of people)"),
        ]:
            add(name, indicator_id, label, number(row[column]), "%", "2022", "NISR RPHC5 Non-Monetary Poverty report", "Table 4.2")
for row in xls_rows(NMP, "Table7_2"):
    name = district_name(row[0])
    if name:
        add(name, "census_mpi_headcount", "Multidimensional poverty headcount, census (% of people)",
            number(row[1]) * 100 if number(row[1]) is not None else None, "%", "2022", "NISR RPHC5 Non-Monetary Poverty report", "Table 7.2")
        add(name, "census_mpi", "Multidimensional Poverty Index (M0), census", number(row[3]), "index 0-1", "2022",
            "NISR RPHC5 Non-Monetary Poverty report", "Table 7.2")

# --- Census 2022 main indicators (district name in column B)
PHC = RAW / "census2022" / "PHC5-2022_Main_Indicators.xlsx"
PHC_TABLES = {
    "Table 1": (2, "census_population", "Resident population", "persons"),
    "Table 25": (2, "census_hh_mobile_phone", "Households with a member owning a mobile phone (%)", "%"),
    "Table 28": (2, "census_medical_insurance", "Population with medical insurance (%)", "%"),
    "Table 59": (3, "census_female_headed_hh", "Households headed by women (%)", "%"),
    "Table 74": (2, "census_electricity", "Households with access to electricity (%)", "%"),
    "Table 21": (5, "census_internet_use_16plus", "Used the internet, age 16+ (%)", "%"),
}
for sheet, (column, indicator_id, label, unit) in PHC_TABLES.items():
    seen = set()
    for row in xlsx_rows(PHC, sheet):
        name = district_name(row[1]) if len(row) > 1 else None
        if name and name not in seen:
            seen.add(name)
            add(name, indicator_id, label, number(row[column]), unit, "2022", "NISR RPHC5 Main Indicators (tables)", sheet)

# --- LFS 2025 district indicators with standard errors (Annex A.1-A.6)
LFS_SE = RAW / "lfs" / "RWLFS2025_STD_ERRORS.xlsx"
for sheet, indicator_id, label in [
    ("UR", "lfs_unemployment_rate", "Unemployment rate (%)"),
    ("LFPR", "lfs_participation_rate", "Labour force participation rate (%)"),
    ("EPR", "lfs_employment_ratio", "Employment-to-population ratio (%)"),
    ("LUUR", "lfs_labour_underutilisation", "Labour underutilisation rate (%)"),
    ("NEET", "lfs_neet_youth", "Youth (16-30) not in employment, education or training (%)"),
    ("POLR", "lfs_outside_labour_force", "Population outside the labour force (%)"),
]:
    seen = set()
    for row in xlsx_rows(LFS_SE, sheet):
        name = district_name(row[0])
        if name and name not in seen:
            seen.add(name)
            add(name, indicator_id, label, number(row[1]), "%", "2025", "NISR LFS 2025 annual (standard errors annex)",
                f"Annex {sheet}", se=number(row[2]), ci_low=number(row[3]), ci_high=number(row[4]))

# --- LFS district time series (Tables 21-25)
LFS = RAW / "lfs" / "RWLFS_Annual_Tables_LFS2025.xlsx"
lfs_series = []
workbook = openpyxl.load_workbook(LFS, read_only=True, data_only=True)
for sheet in workbook.sheetnames:
    if not sheet.startswith(("T21", "T22", "T23", "T24", "T25")):
        continue
    current, years = None, []
    for row in workbook[sheet].iter_rows(values_only=True):
        row = list(row)
        name = district_name(row[0])
        if name:
            current = name
            # The block header repeats the district name twice: first with "Period: Year", then with years
            # (stored as text, e.g. '2017').
            candidates = [number(v) for v in row[1:10] if v is not None]
            if candidates and all(v is not None and 2000 < v < 2100 for v in candidates):
                years = [int(v) for v in candidates]
            continue
        if current and isinstance(row[0], str) and years:
            for year, value in zip(years, row[1:1 + len(years)]):
                value = number(value)
                if value is not None:
                    lfs_series.append({"district": current, "province": PROVINCE[current],
                                       "indicator": re.sub(r"\s+", " ", row[0]).strip(), "year": year, "value": round(value, 3)})
workbook.close()
for item in lfs_series:
    if item["year"] == 2025 and item["indicator"].startswith("Median monthly earnings"):
        add(item["district"], "lfs_median_monthly_earnings", "Median monthly earnings at main job", item["value"], "RWF",
            "2025", "NISR LFS 2025 annual tables", "Tables 21-25")

# --- Population projections: adults 16+ and youth 16-30 (2024 and 2026)
for path in sorted((RAW / "population-projections").glob("*.xlsx")):
    name = next((d for d in DISTRICTS if d.lower() in path.stem.lower()), None)
    if not name:
        continue
    rows = xlsx_rows(path, openpyxl.load_workbook(path, read_only=True).sheetnames[0])
    year_row = next(r for r in rows if sum(isinstance(v, int) and 2020 < v < 2040 for v in r) >= 5)
    columns = {v: i for i, v in enumerate(year_row) if isinstance(v, int) and 2020 < v < 2040}
    totals = defaultdict(float)
    for row in rows:
        age = row[0]
        if isinstance(age, str) and age.strip().endswith("+"):
            age = int(age.strip().rstrip("+"))
        if not isinstance(age, (int, float)):
            continue
        for year in (2024, 2026):
            both, female = number(row[columns[year]]), number(row[columns[year] + 2])
            if age >= 16 and both:
                totals[(year, "adults")] += both
                totals[(year, "women")] += female or 0
            if 16 <= age <= 30 and both:
                totals[(year, "youth")] += both
    for year in (2024, 2026):
        add(name, f"proj_adults_16plus_{year}", f"Projected adults aged 16+ ({year})", totals[(year, "adults")], "persons",
            str(year), "NISR subnational population projections 2023-2032", "single-age tables", status="projection")
        add(name, f"proj_women_16plus_{year}", f"Projected women aged 16+ ({year})", totals[(year, "women")], "persons",
            str(year), "NISR subnational population projections 2023-2032", "single-age tables", status="projection")
        add(name, f"proj_youth_16_30_{year}", f"Projected youth aged 16-30 ({year})", totals[(year, "youth")], "persons",
            str(year), "NISR subnational population projections 2023-2032", "single-age tables", status="projection")

# --- Statistical Yearbook 2025: RSSB active members and Mutuelle registrations by district
yearbook = xlsx_rows(RAW / "other-nisr" / "Rwanda_Statistical_Yearbook_2025.xlsx", "BANKING and FINANCE")
mode = None
for row in yearbook:
    first = row[0]
    if isinstance(first, str) and first.startswith("Table 12.2.3"):
        mode = "rssb"
        continue
    if isinstance(first, str) and first.startswith("Table 12.2.9"):
        mode = "mutuelle"
        continue
    if isinstance(first, str) and first.startswith("Table"):
        mode = None
    name = district_name(first)
    if name and mode == "rssb":
        add(name, "rssb_active_members", "RSSB active (contributing) members", number(row[9]), "persons", "2023/24",
            "NISR Statistical Yearbook 2025 (RSSB data)", "Table 12.2.3")
    if name and mode == "mutuelle":
        add(name, "cbhi_registrations", "Mutuelle de Sante (CBHI) registrations", number(row[8]), "persons", "2024/25",
            "NISR Statistical Yearbook 2025 (RSSB data)", "Table 12.2.9")

# --- Establishment Census 2023 (full count): establishments, informal enterprises, informal employment
EC = RAW / "other-nisr" / "EC_2023_Tables.xlsx"
EC_SOURCE = "NISR Establishment Census 2023 (tables)"
for row in xlsx_rows(EC, "Table 2.1.4"):
    name = district_name(row[0])
    if name:
        add(name, "ec_establishments", "Establishments (all institutional sectors)", number(row[4]), "count", "2023",
            EC_SOURCE, "Table 2.1.4")
for sheet, indicator_id, label in [
    ("Table 3. 10", "ec_informal_enterprises_share", "Informal enterprises (% of enterprises)"),
    ("Table 3. 18", "ec_informal_employment_share", "Informal employment in establishments (% of workers)"),
]:
    for row in xlsx_rows(EC, sheet):
        name, total, informal = district_name(row[0]), number(row[1]), number(row[3])
        if name and total:
            add(name, indicator_id, label, informal / total * 100, "%", "2023", EC_SOURCE,
                f"{sheet} (informal / total counts)", status="calculated")

# --- District values transcribed from other published reports (method and checks: extracts README)
TRANSCRIBED = [
    ("finscope2020_district_banked.csv", "NISR/AFR FinScope 2020 report", "Figure 17 (banked by district)", "2020",
     "% of adults 16+", [
         ("banked_incl_otc_pct", "finscope2020_banked_incl_otc",
          "Adults using bank services, incl. over-the-counter users without an own account (%). Not comparable with finscope_banked"),
     ]),
    ("dhs2025_district_child_nutrition.csv", "NISR Rwanda DHS 2025 final report", "Table D.4", "2025",
     "% of children under 5", [
         ("stunted_pct", "dhs_stunting", "Stunted children under 5 (height-for-age below -2 SD) (%)"),
         ("severely_stunted_pct", "dhs_severe_stunting", "Severely stunted children under 5 (below -3 SD) (%)"),
         ("wasted_pct", "dhs_wasting", "Wasted children under 5 (weight-for-height below -2 SD) (%)"),
         ("underweight_pct", "dhs_underweight", "Underweight children under 5 (weight-for-age below -2 SD) (%)"),
     ]),
    ("cfsva2024_district_food_consumption.csv", "WFP/MINAGRI/NISR CFSVA 2024 report", "Figure 7.3 (chart labels)", "2024",
     "% of households", [
         ("fcs_poor_pct", "cfsva_fcs_poor", "Households with poor food consumption (%)"),
         ("fcs_borderline_pct", "cfsva_fcs_borderline", "Households with borderline food consumption (%)"),
     ]),
    ("cfsva2024_district_natural_hazards.csv", "WFP/MINAGRI/NISR CFSVA 2024 report",
     "Figure 8.4 (bar heights digitized, about +/-1 point)", "2024", "% of households", [
         ("hazard_any_pct", "cfsva_hazard_any", "Households affected by a natural hazard in the 12 months before the survey (%)"),
         ("hazard_drought_pct", "cfsva_hazard_drought", "Households affected mainly by drought (%)"),
         ("hazard_floods_pct", "cfsva_hazard_floods", "Households affected mainly by floods or heavy rains (%)"),
         ("hazard_landslides_pct", "cfsva_hazard_landslides", "Households affected mainly by landslides (%)"),
     ]),
]
for filename, source, table, year, unit, columns in TRANSCRIBED:
    for row in csv.DictReader(open(OUT / filename, encoding="utf-8")):
        name = district_name(row["district"])
        for column, indicator_id, label in columns:
            if name:
                add(name, indicator_id, label, number(row[column]), unit, year, source, table)
        poor, borderline = number(row.get("fcs_poor_pct")), number(row.get("fcs_borderline_pct"))
        if name and poor is not None and borderline is not None:
            add(name, "cfsva_inadequate_food_consumption", "Households with inadequate (poor or borderline) food consumption (%)",
                poor + borderline, unit, year, source, "Figure 7.3 (sum of rounded chart labels)", status="calculated")

# --- Calculated counts (clearly marked): rate x projected population
by_key = {(r["district"], r["indicator_id"]): r for r in records}
for name in DISTRICTS:
    adults = by_key.get((name, "proj_adults_16plus_2024"))
    for rate_id, count_id, label in [
        ("finscope_excluded", "calc_excluded_adults_2024", "Financially excluded adults, approx. (FinScope rate x 2024 projection)"),
        ("finscope_informal_only", "calc_informal_only_adults_2024", "Adults relying on informal finance only, approx."),
    ]:
        rate = by_key.get((name, rate_id))
        if adults and rate:
            add(name, count_id, label, rate["value"] / 100 * adults["value"], "persons", "2024",
                "IMBONIX calculation", f"{rate_id} x proj_adults_16plus_2024", status="calculated")
    rssb = by_key.get((name, "rssb_active_members"))
    if adults and rssb:
        add(name, "calc_rssb_members_per_100_adults", "RSSB active members per 100 adults (formal-sector pension reach proxy)",
            rssb["value"] / adults["value"] * 100, "per 100 adults", "2024", "IMBONIX calculation",
            "rssb_active_members / proj_adults_16plus_2024", status="calculated")
    establishments = by_key.get((name, "ec_establishments"))
    if adults and establishments:
        add(name, "calc_establishments_per_1000_adults", "Establishments per 1,000 adults (2023 census / 2024 adults)",
            establishments["value"] / adults["value"] * 1000, "per 1,000 adults", "2023/24", "IMBONIX calculation",
            "ec_establishments / proj_adults_16plus_2024", status="calculated")

# --- EICV7 VUP thematic tables: benefit amounts, payment channel and timeliness by poverty status
VUP = RAW / "eicv7" / "Final_EICV7_VUP_Thematic_Report_Tables.xlsx"
VUP_TABLES = {"Table 4.2": "Direct Support", "Table 4.5": "Classic Public Works",
              "Table 4.8": "Expanded Public Works", "Table 4.11": "Nutrition-Sensitive Direct Support"}
vup_rows = []
for sheet, component in VUP_TABLES.items():
    block = None
    for row in xlsx_rows(VUP, sheet):
        label = row[1] if len(row) > 1 else None
        if not isinstance(label, str) or not label.strip():
            continue
        label = re.sub(r"\s+", " ", label).strip()
        values = [number(v) for v in row[2:6]]
        if all(v is None for v in values):
            if not label.lower().startswith(("source", "eicv7")):
                block = label
            continue
        if label.lower() == "total" or block is None:
            continue
        vup_rows.append({"component": component, "table": sheet, "block": block, "category": label,
                         "extremely_poor": values[0], "moderately_poor": values[1], "non_poor": values[2],
                         "all_beneficiaries": values[3]})
with (OUT / "vup_benefit_delivery_eicv7.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(vup_rows[0]))
    writer.writeheader()
    writer.writerows(vup_rows)
print(f"VUP delivery rows: {len(vup_rows)}")

# --- Census 2022 thematic reports: persons with disabilities and older people (60+), two groups Direct Support serves
DISABILITY = RAW / "census2022" / "Socio_Economic Status of people with Disability Thematic Report.xlsx"
for row in xlsx_rows(DISABILITY, "Table C.1"):
    name = district_name(row[0])
    if name:
        add(name, "census_disability_prevalence", "Persons with disabilities among residents aged 5+ (%)", number(row[6]),
            "%", "2022", "NISR RPHC5 Persons with Disabilities thematic report (tables)", "Table C.1")
OLDER = RAW / "census2022" / "Socio-Economic status of aged people _Thematic Report.xls"
for row in xls_rows(OLDER, "Table 2"):
    name = district_name(row[1])
    if name:
        add(name, "census_older_people_share", "People aged 60+ among residents (%)", number(row[7]), "%", "2022",
            "NISR RPHC5 Older People thematic report (tables)", "Table 2")
for row in xls_rows(OLDER, "Table 19 "):
    name = district_name(row[1])
    if name:
        add(name, "census_older_people_mobile_phone", "People aged 60+ who own a mobile phone (%)", number(row[4]), "%",
            "2022", "NISR RPHC5 Older People thematic report (tables)", "Table 19")

# --- Write district outputs
fields = ["district", "province", "indicator_id", "indicator", "value", "se", "ci_low", "ci_high", "unit", "year",
          "source", "table", "status"]
records.sort(key=lambda r: (DISTRICTS.index(r["district"]), r["indicator_id"]))
with (OUT / "district_indicators_long.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=fields)
    writer.writeheader()
    writer.writerows(records)
indicator_ids = sorted({r["indicator_id"] for r in records})
wide = {name: {"district": name, "province": PROVINCE[name]} for name in DISTRICTS}
for r in records:
    wide[r["district"]][r["indicator_id"]] = r["value"]
with (OUT / "district_indicators_wide.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=["district", "province", *indicator_ids])
    writer.writeheader()
    writer.writerows(wide.values())
with (OUT / "lfs_district_2017_2025.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=["district", "province", "indicator", "year", "value"])
    writer.writeheader()
    writer.writerows(lfs_series)

# --- Sector outputs: census non-monetary poverty (Table C.1), MPI (Table C.10), population (Tables 96-100)
sectors = {}


def sector_rows(rows, label_col, first_value_col):
    district = None
    for row in rows:
        label = row[label_col] if len(row) > label_col else None
        if not isinstance(label, str) or not label.strip():
            continue
        label = re.sub(r"\s+", " ", label).strip()
        values = [number(v) for v in row[first_value_col:first_value_col + 7]]
        name = district_name(label)
        if name and all(v is None for v in values):
            district = name
            continue
        if name and district is None:  # population tables give the district total on its own row
            district = name
            continue
        if district and any(v is not None for v in values) and not label.lower().startswith(("total", "rwanda", "table")):
            yield district, label, values


# Sector names are keyed case-insensitively (Rulindo's sector appears as "Base" and "BASE").
for district, sector, values in sector_rows(xls_rows(NMP, "Table_C_1"), 0, 1):
    item = sectors.setdefault((district, sector.lower()), {"district": district, "province": PROVINCE[district], "sector": sector})
    item.update({"nonpoor_pct": values[0], "vulnerable_pct": values[1], "moderately_poor_pct": values[2],
                 "severely_poor_pct": values[3]})
for district, sector, values in sector_rows(xls_rows(NMP, "Table_C_10"), 0, 1):
    item = sectors.setdefault((district, sector.lower()), {"district": district, "province": PROVINCE[district], "sector": sector})
    item.update({"mpi_headcount_pct": round(values[0] * 100, 2) if values[0] is not None else None,
                 "mpi_intensity": values[1], "mpi": values[2]})
POPULATION_TABLES = {"Table 96": "Kigali City", "Table 97": "South", "Table 98": "West",
                     "Table 99": "North", "Table 100": "East"}
for sheet, province in POPULATION_TABLES.items():
    district, opened = None, set()
    for row in xlsx_rows(PHC, sheet):
        label = row[1] if len(row) > 1 else None
        if not isinstance(label, str):
            continue
        label = re.sub(r"\s+", " ", label).strip()
        name = district_name(label)
        # A district's total row opens its block the first time its name appears in its province's table.
        # Sectors can share a district's name (sector Nyarugenge; sector Ngoma in Huye; sector Nyanza in
        # Gisagara), so later occurrences are sectors.
        if name and name != district and PROVINCE[name] == province and name not in opened:
            district = name
            opened.add(name)
            continue
        if district and (district, label.lower()) in sectors:
            sectors[(district, label.lower())].update({
                "population_2022": number(row[2]), "population_male_2022": number(row[3]),
                "population_female_2022": number(row[4]),
            })

sector_fields = ["district", "province", "sector", "population_2022", "population_male_2022", "population_female_2022",
                 "nonpoor_pct", "vulnerable_pct", "moderately_poor_pct", "severely_poor_pct",
                 "mpi_headcount_pct", "mpi_intensity", "mpi"]
with (OUT / "sector_census2022.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=sector_fields)
    writer.writeheader()
    for item in sorted(sectors.values(), key=lambda s: (DISTRICTS.index(s["district"]), s["sector"])):
        writer.writerow({k: item.get(k) for k in sector_fields})

# --- Timeline: Rwanda over the years and each district's projected population
# Census totals from 1978 (RPHC5 Table 4); EICV, DHS and LFS rounds from the Statistical Yearbook 2025 trend tables
# (Tables 1.1 to 1.3); province poverty in 2016/17 (modelled on the EICV7 method) and 2023/24 (EICV7 poverty profile,
# Table 5.1); and each district's projected population for every year from 2023 to 2032.
timeline = []


def add_point(series_id, label, unit, period, value, area, source, table, status="observed", note=""):
    """One published value for one period: a single year (2022) or a survey period (2023/24)."""
    if value is None:
        return
    first, _, last = str(period).partition("/")
    start = int(first)
    end = int(first[:2] + last) if last else start
    timeline.append({"series_id": series_id, "label": label, "unit": unit, "period": str(period), "start": start,
                     "end": end, "area": area, "value": round(value, 4), "source": source, "table": table,
                     "status": status, "note": note})


def clean_label(value):
    return re.sub(r"\s+", " ", value).strip() if isinstance(value, str) else ""


for row in xlsx_rows(PHC, "Table 4"):
    year = row[1] if len(row) > 1 else None
    if isinstance(year, (int, float)) and 1970 < year < 2030:
        add_point("census_population", "Population", "persons", int(year), number(row[2]), "Rwanda",
                  "NISR RPHC5 Main Indicators (tables)", "Table 4")

YEARBOOK_TRENDS = [
    # Table 1.1, household indicators by EICV round: the indicator name is in column B, values from column C.
    ("EICV", lambda row: clean_label(row[1] if len(row) > 1 else None) == "Indicator Name", 1, {
        r"^Average household size": ("household_size", "Average household size", "persons"),
        r"ever attended school": ("ever_attended_school", "People aged 6 and over who ever attended school (%)", "%"),
        r"Living in Umudugudu": ("umudugudu", "Households living in umudugudu (%)", "%"),
        r"metal \(iron\) sheet roof": ("metal_roof", "Households with an iron sheet roof (%)", "%"),
        r"cement floor": ("cement_floor", "Households with a cement floor (%)", "%"),
        r"electricity as main source of lighting": ("electricity_lighting", "Households lit mainly by electricity (%)", "%"),
        r"firewood as main cooking fuel": ("firewood_cooking", "Households cooking mainly with firewood (%)", "%"),
        r"improved drinking water": ("improved_water", "Households using improved drinking water (%)", "%"),
        r"improved sanitation": ("improved_sanitation", "Households using improved sanitation (%)", "%"),
        r"internet at home": ("internet_home", "Households with internet access at home (%)", "%"),
        r"Owning mobile phone": ("mobile_phone", "Households owning a mobile phone (%)", "%"),
        r"reach a health center": ("minutes_to_health_centre", "Average time to reach a health centre (minutes)", "minutes"),
        r"health insurance": ("health_insurance", "People with health insurance (%)", "%"),
    }, "Table 1.1", ""),
    # Table 1.2, demographic and health indicators by DHS round: the indicator name is in column A.
    ("DHS", lambda row: clean_label(row[0] if row else None) == "Indicators" and str(row[1]).strip() == "1992", 0, {
        r"^Total fertility Rate": ("fertility_rate", "Total fertility rate (children per woman)", "children"),
        r"^The use of Modern Contraceptive": ("modern_contraception", "Married women using modern contraception (%)", "%"),
        r"^Vaccination": ("vaccination", "Children vaccinated (%)", "%"),
        r"^Infant mortality": ("infant_mortality", "Infant mortality (per 1,000 live births)", "per 1,000"),
        r"^Child mortality": ("under_five_mortality", "Under five mortality (per 1,000 live births)", "per 1,000"),
        r"^Stunted": ("stunting", "Children under five who are stunted (%)", "%"),
        r"^Wasted": ("wasting", "Children under five who are wasted (%)", "%"),
        r"^Underweight": ("underweight", "Children under five who are underweight (%)", "%"),
        r"^Maternal mortality ratio": ("maternal_mortality", "Maternal mortality (per 100,000 live births)", "per 100,000"),
        r"^Assistance during delivery": ("assisted_delivery", "Births with assistance during delivery (%)", "%"),
    }, "Table 1.2", "Some DHS rates refer to the years before the survey."),
    # Table 1.3, national labour force indicators by year: the indicator name is in column A.
    ("LFS", lambda row: clean_label(row[0] if row else None) == "Indicators" and str(row[1]).strip() == "2019", 0, {
        r"^Labour force participation rate": ("labour_force_participation", "Labour force participation rate (%)", "%"),
        r"^Employment to population ratio": ("employment_to_population", "Employment to population ratio (%)", "%"),
        r"^Unemployment rate \(%\)$": ("unemployment", "Unemployment rate (%)", "%"),
    }, "Table 1.3", ""),
]
YEARBOOK_FILE = RAW / "other-nisr" / "Rwanda_Statistical_Yearbook_2025.xlsx"
trend_rows = xlsx_rows(YEARBOOK_FILE, openpyxl.load_workbook(YEARBOOK_FILE, read_only=True).sheetnames[0])
for survey, is_header, label_col, patterns, table, note in YEARBOOK_TRENDS:
    header = next(i for i, row in enumerate(trend_rows) if len(row) > 1 and is_header(row))
    periods = {col: str(value).strip() for col, value in enumerate(trend_rows[header]) if col > label_col and value is not None}
    for row in trend_rows[header + 1:]:
        label = clean_label(row[label_col] if len(row) > label_col else None)
        if label.startswith(("Source", "Table")):
            break
        match = next((spec for pattern, spec in patterns.items() if re.search(pattern, label)), None)
        if not match:
            continue
        series_id, series_label, unit = match
        for col, period in periods.items():
            add_point(series_id, series_label, unit, period, number(row[col]) if col < len(row) else None, "Rwanda",
                      f"NISR Statistical Yearbook 2025 ({survey} rounds)", table, note=note)

POVERTY_PROFILE = RAW / "eicv7" / "EICV7_Tables_Rwanda_Poverty_Profile.xlsx"
profile_sheet = next(s for s in openpyxl.load_workbook(POVERTY_PROFILE, read_only=True).sheetnames if s.strip() == "Table 5.1.")
for row in xlsx_rows(POVERTY_PROFILE, profile_sheet):
    area = clean_label(row[0] if row else None)
    if area == "Area of residence":
        break
    if area not in ("Rwanda", "Kigali City", "South", "West", "North", "East"):
        continue
    for series_id, label, now_col, before_col in (("poverty_rate", "Poverty rate (%)", 1, 3),
                                                  ("extreme_poverty_rate", "Extreme poverty rate (%)", 6, 8)):
        add_point(series_id, label, "%", "2016/17", number(row[before_col]), area,
                  "NISR EICV7 Rwanda Poverty Profile (tables)", "Table 5.1", status="model_estimate",
                  note="2016/17 is NISR's estimate on the EICV7 method, so it compares directly with 2023/24.")
        add_point(series_id, label, "%", "2023/24", number(row[now_col]), area,
                  "NISR EICV7 Rwanda Poverty Profile (tables)", "Table 5.1")

projected = []
for path in sorted((RAW / "population-projections").glob("*.xlsx")):
    name = next((d for d in DISTRICTS if d.lower() in path.stem.lower()), None)
    if not name:
        continue
    rows = xlsx_rows(path, openpyxl.load_workbook(path, read_only=True).sheetnames[0])
    year_row = next(r for r in rows if sum(isinstance(v, int) and 2020 < v < 2040 for v in r) >= 5)
    columns = {v: i for i, v in enumerate(year_row) if isinstance(v, int) and 2020 < v < 2040}
    total = next(r for r in rows if isinstance(r[0], str) and r[0].strip().lower() == "total")
    for year, col in sorted(columns.items()):
        projected.append({"district": name, "province": PROVINCE[name], "year": year, "population": number(total[col])})
projected.sort(key=lambda r: (DISTRICTS.index(r["district"]), r["year"]))

timeline_fields = ["series_id", "label", "unit", "period", "start", "end", "area", "value", "source", "table", "status",
                   "note"]
with (OUT / "timeline_national.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=timeline_fields, lineterminator="\n")
    writer.writeheader()
    writer.writerows(timeline)
with (OUT / "district_population_2023_2032.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=["district", "province", "year", "population"], lineterminator="\n")
    writer.writeheader()
    writer.writerows(projected)

print(f"district records: {len(records)} ({len(indicator_ids)} indicators)")
print(f"LFS series rows: {len(lfs_series)}")
print(f"sectors: {len(sectors)}; with population: {sum(1 for s in sectors.values() if s.get('population_2022'))}")
print(f"timeline points: {len(timeline)}; projected district years: {len(projected)}")
